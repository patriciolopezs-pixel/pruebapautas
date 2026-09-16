import streamlit as st
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, Border, Side, PatternFill
from openpyxl.utils import get_column_letter
import io
import datetime
import re

# Configuración de la página web
st.set_page_config(page_title="Portal de Pautas de Supervisión - RedSalud", layout="centered")

# --- ESTILOS CSS CORPORATIVOS REDSALUD ---
st.markdown("""
<style>
    .stApp {
        background-color: #F4F7F6 !important;
    }
    
    h1, h2, h3, h4 {
        color: #00205B !important;
        font-family: 'Segoe UI', Tahoma, sans-serif !important;
        font-weight: 700 !important;
    }
    
    .stMarkdown, label, p, span {
        color: #00205B !important;
    }

    input, select, textarea, div[data-baseweb="select"] {
        background-color: #FFFFFF !important;
        color: #00205B !important;
        -webkit-text-fill-color: #00205B !important;
        border: 1px solid #CBD5E1 !important;
        border-radius: 8px !important;
        text-transform: uppercase !important;
    }

    button[kind="primary"] {
        background-color: #00828A !important;
        border: none !important;
        border-radius: 8px !important;
    }
    button[kind="primary"] p {
        color: #FFFFFF !important;
        font-weight: bold !important;
    }

    button[kind="secondary"], .stDownloadButton button {
        background-color: #00205B !important;
        border: none !important;
        border-radius: 8px !important;
    }
    button[kind="secondary"] p, .stDownloadButton button p {
        color: #FFFFFF !important;
        font-weight: bold !important;
    }

    .card-pauta {
        background-color: #FFFFFF;
        border: 1px solid #E2E8F0;
        border-top: 6px solid #00828A;
        border-radius: 12px;
        padding: 20px;
        margin-bottom: 20px;
        box-shadow: 0 4px 12px rgba(0, 32, 91, 0.08);
    }
</style>
""", unsafe_allow_html=True)


# --- ESTADO DE LA SESIÓN ---
if 'pagina_activa' not in st.session_state:
    st.session_state.pagina_activa = "inicio"

# 1. Pauta: Pausa Dental (18)
if 'pautas_data' not in st.session_state:
    st.session_state.pautas_data = {i: None for i in range(1, 19)}
if 'pauta_actual' not in st.session_state:
    st.session_state.pauta_actual = 1
if 'ultimo_centro' not in st.session_state:
    st.session_state.ultimo_centro = ""

# 2. Pauta: Higiene de Manos (12)
if 'higiene_data' not in st.session_state:
    st.session_state.higiene_data = {i: None for i in range(1, 13)}
if 'higiene_actual' not in st.session_state:
    st.session_state.higiene_actual = 1
if 'centro_higiene' not in st.session_state:
    st.session_state.centro_higiene = "CD LA REINA"
if 'mes_higiene' not in st.session_state:
    st.session_state.mes_higiene = "AGOSTO 2026"
if 'evaluador_higiene' not in st.session_state:
    st.session_state.evaluador_higiene = ""
if 'responsable_higiene' not in st.session_state:
    st.session_state.responsable_higiene = ""

# 3. Pauta: REG 1.2 AO (18)
if 'reg12_data' not in st.session_state:
    st.session_state.reg12_data = {i: None for i in range(1, 19)}
if 'reg12_actual' not in st.session_state:
    st.session_state.reg12_actual = 1
if 'centro_reg12' not in st.session_state:
    st.session_state.centro_reg12 = "CD LA REINA"
if 'ano_reg12' not in st.session_state:
    st.session_state.ano_reg12 = "2026"
if 'mes_reg12' not in st.session_state:
    st.session_state.mes_reg12 = "JULIO"
if 'grupo_reg12' not in st.session_state:
    st.session_state.grupo_reg12 = "DENTAL"
if 'especialidad_reg12' not in st.session_state:
    st.session_state.especialidad_reg12 = "ODONTOLOGÍA GENERAL"


def parse_fecha(fecha_str):
    if fecha_str:
        try:
            return datetime.datetime.strptime(str(fecha_str).strip(), "%d-%m-%Y").date()
        except Exception:
            try:
                return pd.to_datetime(fecha_str, dayfirst=True).date()
            except Exception:
                pass
    return datetime.date.today()


def clean_text_spaces(text):
    if not text or pd.isna(text):
        return ""
    cleaned = re.sub(r'[^a-zA-ZáéíóúÁÉÍÓÚñÑ\s\.]', '', str(text))
    return re.sub(r'\s+', ' ', cleaned).strip().upper()


# --- LÓGICA CARGA MASIVA PAUSA DENTAL ---
def procesar_df_masivo(df, fecha_sup_default):
    df.columns = [str(c).strip().lower() for c in df.columns]
    
    def get_val(row, candidates, default=""):
        for cand in candidates:
            for col in df.columns:
                if cand in col:
                    val = str(row[col]).strip()
                    return val if val != 'nan' and val != 'None' else default
        return default

    count = 0
    for i, row in df.iterrows():
        if count >= 18:
            break
        num_pauta = count + 1
        
        centro_val = clean_text_spaces(get_val(row, ['centro']))
        rut_val = re.sub(r'[^0-9kK\-]', '', get_val(row, ['rut paciente', 'rut_paciente'])).strip().upper()
        
        nom_prof = get_val(row, ['nombre profesional', 'nombre realizador', 'nombre persona'])
        ape_prof = get_val(row, ['apellidos profesional', 'apellido realizador', 'apellidos persona', 'apellido persona'])
        
        nombre_clean = clean_text_spaces(nom_prof)
        apellido_clean = clean_text_spaces(ape_prof)
        
        fecha_atencion_raw = get_val(row, ['fecha de ejecución', 'fecha atencion', 'fecha_atencion', 'fecha'])
        try:
            fecha_atencion_clean = pd.to_datetime(fecha_atencion_raw, dayfirst=True).strftime("%d-%m-%Y")
        except Exception:
            fecha_atencion_clean = fecha_sup_default.strftime("%d-%m-%Y")
            
        serv_raw = get_val(row, ['servicio']).lower()
        if 'pabell' in serv_raw or 'pd' in serv_raw:
            serv_clean = "Pabellón de Cirugía menor Dental (PD)"
        elif 'imagen' in serv_raw or 'rx' in serv_raw:
            serv_clean = "Imagenología Dental (RX)"
        else:
            serv_clean = "Sala de Procedimiento Dental (BD)"
            
        exo_raw = get_val(row, ['exodoncia', 'prestación corresponde']).upper()
        exo_clean = "SI" if "SI" in exo_raw else "NO"

        cumple_existente = "SI"
        if st.session_state.pautas_data[num_pauta] is not None:
            cumple_existente = st.session_state.pautas_data[num_pauta].get("cumple", "SI")

        st.session_state.pautas_data[num_pauta] = {
            "centro": centro_val if centro_val else st.session_state.ultimo_centro,
            "fecha_sup": fecha_sup_default.strftime("%d-%m-%Y"),
            "nombre": nombre_clean,
            "apellido": apellido_clean,
            "rut": rut_val,
            "fecha_atencion": fecha_atencion_clean,
            "servicio": serv_clean,
            "exodoncia": exo_clean,
            "cumple": cumple_existente
        }
        if centro_val:
            st.session_state.ultimo_centro = centro_val
        count += 1

    return count


# --- LÓGICA CARGA MASIVA REG 1.2 AO (ADAPTADA PARA TEXTO CONCATENADO DE LA GEMA) ---
def procesar_texto_reg12(texto):
    if not texto or not texto.strip():
        return 0

    raw_text = texto.strip()

    # Expresión regular que detecta registros en texto concatenado
    record_pattern = re.compile(
        r'(\d{7,8}-[\dkK])'                        # 1: RUT Paciente
        r'(.*?)'                                   # 2: Nombre Paciente
        r'(\d{1,2}/\d{1,2}/\d{2,4})'               # 3: Fecha Diagnóstico
        r'(.*?)'                                   # 4: Bloque Nombre/RUT Profesional
        r'((?:SI|NO){10})',                        # 5: Secuencia de 10 respuestas SI/NO
        re.IGNORECASE | re.DOTALL
    )

    matches = list(record_pattern.finditer(raw_text))
    count = 0

    if len(matches) > 0:
        for m in matches:
            if count >= 18:
                break
            num_pauta = count + 1

            rut_pac = m.group(1).strip().upper()
            nom_pac = clean_text_spaces(m.group(2))
            fecha_diag_raw = m.group(3).strip()

            try:
                fecha_diag = pd.to_datetime(fecha_diag_raw, dayfirst=True).strftime("%d-%m-%Y")
            except Exception:
                fecha_diag = fecha_diag_raw

            prof_block = m.group(4).strip()
            sino_block = m.group(5).strip().upper()

            nom_prof = ""
            rut_prof = ""

            if "NO REGISTRADO" in prof_block.upper():
                nom_prof = "NO REGISTRADO"
                rut_prof = "NO REGISTRADO"
            else:
                rut_prof_match = re.search(r'(\d{7,8}-[\dkK])', prof_block)
                if rut_prof_match:
                    rut_prof = rut_prof_match.group(1).upper()
                    nom_prof = clean_text_spaces(prof_block[:rut_prof_match.start()])
                else:
                    nom_prof = clean_text_spaces(prof_block)
                    rut_prof = ""

            sino_list = re.findall(r'(SI|NO)', sino_block)

            def get_sino_item(idx, default="SI"):
                if idx < len(sino_list):
                    return sino_list[idx]
                return default

            st.session_state.reg12_data[num_pauta] = {
                "rut_pac": rut_pac,
                "nom_pac": nom_pac,
                "fecha_diag": fecha_diag,
                "nom_prof": nom_prof if nom_prof else "NO REGISTRADO",
                "rut_prof": rut_prof if rut_prof else "NO REGISTRADO",
                "motivo": get_sino_item(0),
                "patologias": get_sino_item(1),
                "medicamentos": get_sino_item(2),
                "alergias": get_sino_item(3),
                "extraoral": get_sino_item(4),
                "intraoral": get_sino_item(5),
                "diagnostico": get_sino_item(6),
                "plan": get_sino_item(7),
                "pronostico": get_sino_item(8),
                "cumple": get_sino_item(9)
            }
            count += 1
        return count

    # Fallback para tablas con tabulaciones desde Excel
    lines = [line.strip() for line in raw_text.split('\n') if line.strip()]
    start_idx = 1 if ('RUT' in lines[0].upper() or 'PACIENTE' in lines[0].upper()) else 0

    for line in lines[start_idx:]:
        if count >= 18:
            break
        num_pauta = count + 1
        parts = [p.strip() for p in re.split(r'\t|;|\s{2,}', line) if p.strip()]
        if len(parts) < 3:
            continue

        ruts = [p for p in parts if re.match(r'^\d{7,8}-[\dkK]$', p)]
        rut_pac = ruts[0] if len(ruts) >= 1 else ""
        rut_prof = ruts[1] if len(ruts) >= 2 else ""

        dates = [p for p in parts if re.search(r'\d{1,2}[/-]\d{1,2}[/-]\d{2,4}', p)]
        fecha_diag = dates[0] if dates else datetime.date.today().strftime("%d-%m-%Y")

        st.session_state.reg12_data[num_pauta] = {
            "rut_pac": re.sub(r'[^0-9kK\-]', '', rut_pac).strip().upper(),
            "nom_pac": clean_text_spaces(parts[1]) if len(parts) > 1 else "",
            "fecha_diag": fecha_diag,
            "nom_prof": clean_text_spaces(parts[3]) if len(parts) > 3 else "NO REGISTRADO",
            "rut_prof": re.sub(r'[^0-9kK\-]', '', rut_prof).strip().upper() if rut_prof else "NO REGISTRADO",
            "motivo": "SI", "patologias": "SI", "medicamentos": "SI", "alergias": "SI",
            "extraoral": "SI", "intraoral": "SI", "diagnostico": "SI", "plan": "SI",
            "pronostico": "SI", "cumple": "SI"
        }
        count += 1

    return count


# --- EXCEL 1: CONSOLIDADO PAUSA DENTAL ---
def generar_excel_pausa_dental(pautas_dict):
    wb = Workbook()
    ws = wb.active
    ws.title = "Consolidado Pautas"

    thin_border = Border(left=Side(style='thin', color='000000'), right=Side(style='thin', color='000000'), top=Side(style='thin', color='000000'), bottom=Side(style='thin', color='000000'))
    center_aligned = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_aligned = Alignment(horizontal="left", vertical="center", wrap_text=True)
    
    bold_font_white = Font(bold=True, color="FFFFFF")
    bold_font_navy = Font(bold=True, color="00205B")
    navy_header_fill = PatternFill(start_color="00205B", end_color="00205B", fill_type="solid")
    teal_sub_fill = PatternFill(start_color="00828A", end_color="00828A", fill_type="solid")
    soft_teal_fill = PatternFill(start_color="E6F7F5", end_color="E6F7F5", fill_type="solid")

    ws.merge_cells('A1:AK1')
    ws['A1'] = "PAUTA DE SUPERVISIÓN CUMPLIMIENTO DE PAUSA DE SEGURIDAD DENTAL EN BOX DENTAL, PABELLÓN DE CIRUGÍA MENOR DENTAL E IMAGENOLOGÍA DENTAL (GCL 2.1 AO)"
    ws['A1'].font = bold_font_white
    ws['A1'].fill = navy_header_fill
    ws['A1'].alignment = center_aligned

    ws['A2'] = "Indicaciones llenado pauta"
    ws.merge_cells('B2:AK2')
    ws['B2'] = "Marque √ SI cumple, Marque X NO cumple, o No Aplica (si corresponde). En ítem cumple registre SI o NO. Registre en observaciones motivo incumplimiento."
    ws['A2'].font = bold_font_navy
    ws['B2'].alignment = center_aligned

    etiquetas = ["Centro", "Fecha de Supervisión", "Nombre de la persona supervisada", "Apellido(s) de la persona supervisada", "RUT del paciente", "Fecha de Atención supervisada", "Servicio Clínico donde se realizó el procedimiento, ya sea Sala de Procedimiento Dental (BD), Pabellón de Cirugía menor Dental (PD), e Imagenología Dental (RX)", "Procedimiento corresponde a Exodoncia (SI, NO)"]

    for i, etiqueta in enumerate(etiquetas, start=3):
        ws.cell(row=i, column=1, value=etiqueta).font = bold_font_navy
        ws.cell(row=i, column=1).alignment = left_aligned
        ws.cell(row=i, column=1).fill = soft_teal_fill

    ws.column_dimensions['A'].width = 50

    ws.cell(row=11, column=1, value="N° DE PAUTA").font = bold_font_white
    ws.cell(row=11, column=1).fill = teal_sub_fill
    ws.cell(row=12, column=1, value="CRITERIOS A EVALUAR").font = bold_font_white
    ws.cell(row=12, column=1).fill = teal_sub_fill
    ws.cell(row=13, column=1, value="Se constata Pausa de Seguridad Dental realizada y registrada en Ficha Clínica.").alignment = left_aligned
    ws.cell(row=14, column=1, value="Cumple (SI/NO)").font = bold_font_navy
    ws.cell(row=14, column=1).fill = soft_teal_fill
    ws.cell(row=15, column=1, value="Total Cumple").font = bold_font_navy
    ws.cell(row=15, column=1).fill = soft_teal_fill

    total_cumple = 0
    total_no_cumple = 0

    for idx in range(18):
        num_pauta = idx + 1
        col_start = 2 + (idx * 2)
        col_end = col_start + 1

        pauta_data = pautas_dict.get(num_pauta) or {}
        campos = ["centro", "fecha_sup", "nombre", "apellido", "rut", "fecha_atencion", "servicio", "exodoncia"]
        for row_idx, campo in enumerate(campos, start=3):
            ws.merge_cells(start_row=row_idx, start_column=col_start, end_row=row_idx, end_column=col_end)
            ws.cell(row=row_idx, column=col_start, value=pauta_data.get(campo, "")).alignment = center_aligned

        ws.merge_cells(start_row=11, start_column=col_start, end_row=11, end_column=col_end)
        ws.cell(row=11, column=col_start, value=num_pauta).alignment = center_aligned
        ws.cell(row=11, column=col_start).font = bold_font_navy
        ws.cell(row=11, column=col_start).fill = soft_teal_fill

        ws.cell(row=12, column=col_start, value="SI").alignment = center_aligned
        ws.cell(row=12, column=col_start).font = bold_font_navy
        ws.cell(row=12, column=col_end, value="NO").alignment = center_aligned
        ws.cell(row=12, column=col_end).font = bold_font_navy

        cumple = pauta_data.get("cumple", "")
        if cumple == "SI":
            ws.cell(row=13, column=col_start, value="√").alignment = center_aligned
            ws.merge_cells(start_row=14, start_column=col_start, end_row=14, end_column=col_end)
            ws.cell(row=14, column=col_start, value="SI").alignment = center_aligned
            total_cumple += 1
        elif cumple == "NO":
            ws.cell(row=13, column=col_end, value="X").alignment = center_aligned
            ws.merge_cells(start_row=14, start_column=col_start, end_row=14, end_column=col_end)
            ws.cell(row=14, column=col_start, value="NO").alignment = center_aligned
            total_no_cumple += 1
        else:
            ws.merge_cells(start_row=14, start_column=col_start, end_row=14, end_column=col_end)

    ws.merge_cells('B15:F15')
    ws['B15'] = total_cumple
    ws['B15'].alignment = center_aligned

    ws.merge_cells('G15:J15')
    ws['G15'] = "Total No Cumple"
    ws['G15'].font = bold_font_navy
    ws['G15'].alignment = center_aligned

    ws.merge_cells('K15:N15')
    ws['K15'] = total_no_cumple
    ws['K15'].alignment = center_aligned

    ws.merge_cells('O15:R15')
    ws['O15'] = "% Cumplimiento"
    ws['O15'].font = bold_font_navy
    ws['O15'].alignment = center_aligned

    completadas = sum(1 for v in pautas_dict.values() if v is not None)
    porcentaje = f"{(total_cumple/18)*100:.1f}%" if completadas == 18 else f"{(total_cumple/completadas)*100:.1f}%" if completadas > 0 else "-"
    ws.merge_cells('S15:V15')
    ws['S15'] = porcentaje
    ws['S15'].alignment = center_aligned

    ws.merge_cells('A16:Q20')
    ws['A16'] = "Observaciones:"
    ws['A16'].font = bold_font_navy
    ws['A16'].alignment = Alignment(horizontal="left", vertical="top")

    ws.merge_cells('R16:U18')
    ws.merge_cells('R19:U20')
    ws['R19'] = "Nombre o Timbre\ndel responsable de\naplicar la pauta"
    ws['R19'].font = bold_font_navy
    ws['R19'].alignment = center_aligned

    ws.row_dimensions[19].height = 20
    ws.row_dimensions[20].height = 20

    for r in range(1, 16):
        for c in range(1, 38):
            ws.cell(row=r, column=c).border = thin_border

    for r in range(16, 21):
        for c in range(1, 22):
            ws.cell(row=r, column=c).border = thin_border

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


# --- EXCEL 2: CONSOLIDADO HIGIENE DE MANOS (GCL 1.2) ---
def generar_excel_higiene_manos(higiene_dict, centro, mes, responsable):
    wb = Workbook()
    ws = wb.active
    ws.title = "Higiene de Manos"

    thin_border = Border(left=Side(style='thin', color='000000'), right=Side(style='thin', color='000000'), top=Side(style='thin', color='000000'), bottom=Side(style='thin', color='000000'))
    center_aligned = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_aligned = Alignment(horizontal="left", vertical="center", wrap_text=True)
    
    bold_font_white = Font(bold=True, color="FFFFFF")
    bold_font_navy = Font(bold=True, color="00205B")
    navy_header_fill = PatternFill(start_color="00205B", end_color="00205B", fill_type="solid")
    soft_teal_fill = PatternFill(start_color="E6F7F5", end_color="E6F7F5", fill_type="solid")

    ws.merge_cells('A1:M1')
    ws['A1'] = "GCL 1.2 PAUTA SUPERVISIÓN DE HIGIENE DE MANOS - ÁREA DENTAL"
    ws['A1'].font = bold_font_white
    ws['A1'].fill = navy_header_fill
    ws['A1'].alignment = center_aligned

    ws['A3'] = "Centro Dental:"
    ws['A3'].font = bold_font_navy
    ws['B3'] = centro
    ws['B3'].alignment = left_aligned

    ws['A4'] = "Mes:"
    ws['A4'].font = bold_font_navy
    ws['B4'] = mes
    ws['B4'].alignment = left_aligned

    ws.cell(row=5, column=1, value="Número correlativo").font = bold_font_navy
    ws.cell(row=6, column=1, value="Fecha de la evaluación").font = bold_font_navy
    ws.cell(row=7, column=1, value="Nombre del evaluador").font = bold_font_navy
    ws.cell(row=8, column=1, value="Nombre del evaluado").font = bold_font_navy

    for r in range(5, 9):
        ws.cell(row=r, column=1).fill = soft_teal_fill
        ws.cell(row=r, column=1).alignment = left_aligned

    ws.column_dimensions['A'].width = 38

    ws.cell(row=10, column=1, value="CRITERIOS A EVALUAR").font = bold_font_navy
    ws.cell(row=10, column=1).fill = soft_teal_fill
    ws.cell(row=11, column=1, value="N° oportunidad evaluada").font = bold_font_navy
    ws.cell(row=12, column=1, value="Se realiza higiene de manos según oportunidad").font = bold_font_navy

    total_cumple = 0
    total_aplicadas = 0

    for idx in range(12):
        num = idx + 1
        col = idx + 2
        ws.column_dimensions[get_column_letter(col)].width = 15
        
        ws.cell(row=5, column=col, value=num).alignment = center_aligned
        ws.cell(row=5, column=col).font = bold_font_navy

        data = higiene_dict.get(num) or {}
        ws.cell(row=6, column=col, value=data.get("fecha_eval", "")).alignment = center_aligned
        ws.cell(row=7, column=col, value=data.get("evaluador", "")).alignment = center_aligned
        ws.cell(row=8, column=col, value=data.get("evaluado", "")).alignment = center_aligned
        ws.cell(row=11, column=col, value=data.get("oportunidad", "")).alignment = center_aligned
        
        cumple = data.get("cumple", "")
        if cumple:
            ws.cell(row=12, column=col, value=cumple).alignment = center_aligned
            total_aplicadas += 1
            if cumple == "SI":
                total_cumple += 1

    ws.merge_cells('A14:E14')
    ws['A14'] = "Oportunidades de lavado de manos"
    ws['A14'].font = bold_font_navy
    ws['A14'].fill = soft_teal_fill

    leyendas = ["1. Antes del contacto con el paciente", "2. Antes de una técnica aséptica", "3. Después de la exposición a fluidos corporales o manejo de fluidos contaminados", "4. Después del contacto con el paciente", "5. Después de tener contacto con la zona alrededor del paciente"]
    for i, ley in enumerate(leyendas, start=15):
        ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=5)
        ws.cell(row=i, column=1, value=ley).alignment = left_aligned

    ws.cell(row=21, column=1, value="Total de pautas que cumplen criterios").font = bold_font_navy
    ws.cell(row=21, column=2, value=total_cumple).alignment = center_aligned

    ws.cell(row=22, column=1, value="Total de pautas aplicadas").font = bold_font_navy
    ws.cell(row=22, column=2, value=total_aplicadas).alignment = center_aligned

    pct = f"{(total_cumple/total_aplicadas)*100:.0f}%" if total_aplicadas > 0 else "0%"
    ws.cell(row=23, column=1, value="Porcentaje de cumplimiento").font = bold_font_navy
    ws.cell(row=23, column=2, value=pct).alignment = center_aligned

    ws.cell(row=25, column=1, value="Nombre de responsable del indicador").font = bold_font_navy
    ws.merge_cells('B25:E25')
    ws['B25'] = responsable
    ws['B25'].alignment = left_aligned

    for c in range(1, 14):
        ws.cell(row=1, column=c).border = thin_border

    ws.cell(row=3, column=1).border = thin_border
    ws.cell(row=3, column=2).border = thin_border
    ws.cell(row=4, column=1).border = thin_border
    ws.cell(row=4, column=2).border = thin_border

    for r in range(5, 9):
        for c in range(1, 14):
            ws.cell(row=r, column=c).border = thin_border

    for r in range(10, 13):
        for c in range(1, 14):
            ws.cell(row=r, column=c).border = thin_border

    for r in range(14, 20):
        for c in range(1, 6):
            ws.cell(row=r, column=c).border = thin_border

    for r in range(21, 24):
        for c in range(1, 3):
            ws.cell(row=r, column=c).border = thin_border

    ws.cell(row=25, column=1).border = thin_border
    for c in range(2, 6):
        ws.cell(row=25, column=c).border = thin_border

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


# --- EXCEL 3: CONSOLIDADO REG 1.2 AO (REGISTROS MÍNIMOS EN FICHA CLÍNICA) ---
def generar_excel_reg12(reg_dict, centro, ano, mes, grupo, especialidad):
    wb = Workbook()
    ws = wb.active
    ws.title = "Consolidado REG 1.2"

    thin_border = Border(left=Side(style='thin', color='000000'), right=Side(style='thin', color='000000'), top=Side(style='thin', color='000000'), bottom=Side(style='thin', color='000000'))
    center_aligned = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_aligned = Alignment(horizontal="left", vertical="center", wrap_text=True)
    
    bold_font_white = Font(bold=True, color="FFFFFF")
    bold_font_navy = Font(bold=True, color="00205B")
    navy_header_fill = PatternFill(start_color="00205B", end_color="00205B", fill_type="solid")
    teal_sub_fill = PatternFill(start_color="00828A", end_color="00828A", fill_type="solid")
    soft_teal_fill = PatternFill(start_color="E6F7F5", end_color="E6F7F5", fill_type="solid")

    ws.merge_cells('A1:AK1')
    ws['A1'] = "PAUTA DE SUPERVISIÓN REGISTROS MINIMOS EN FICHA CLÍNICA DENTAL (REG 1.2 AO)"
    ws['A1'].font = bold_font_white
    ws['A1'].fill = navy_header_fill
    ws['A1'].alignment = center_aligned

    ws['A2'] = "Indicaciones llenado pauta"
    ws.merge_cells('B2:AK2')
    ws['B2'] = "Marque √ SI cumple, Marque X NO cumple. En item cumple registre SI o NO. Pauta Dicotómica Registre en observaciones motivo incumplimiento."
    ws['A2'].font = bold_font_navy
    ws['B2'].alignment = center_aligned

    metadata_labels = [
        ("Centro", centro),
        ("Año", ano),
        ("Mes", mes),
        ("Grupo", grupo),
        ("Especialidad", especialidad),
        ("RUT Paciente", "rut_pac"),
        ("Nombre Paciente", "nom_pac"),
        ("Fecha", "fecha_diag"),
        ("Nombre Profesional", "nom_prof"),
        ("RUT Profesional", "rut_prof")
    ]

    for idx, (label, val_key) in enumerate(metadata_labels, start=3):
        ws.cell(row=idx, column=1, value=label).font = bold_font_navy
        ws.cell(row=idx, column=1).fill = soft_teal_fill
        ws.cell(row=idx, column=1).alignment = left_aligned

    ws.column_dimensions['A'].width = 45

    ws.cell(row=13, column=1, value="N° DE PAUTA").font = bold_font_white
    ws.cell(row=13, column=1).fill = teal_sub_fill

    ws.cell(row=14, column=1, value="CRITERIOS A EVALUAR").font = bold_font_white
    ws.cell(row=14, column=1).fill = teal_sub_fill

    rows_structure = [
        (15, "1. HISTORIA CLÍNICA", True),
        (16, "1.1 Anamnesis", True),
        (17, "Motivo de Consulta", "motivo"),
        (18, "Patologías", "patologias"),
        (19, "Medicamentos", "medicamentos"),
        (20, "Alergias", "alergias"),
        (21, "1.2. Examen Extraoral", "extraoral"),
        (22, "1.3. Examen Intraoral", "intraoral"),
        (23, "2. DIAGNÓSTICO Y PLAN DE TRATAMIENTO", True),
        (24, "2.1. Diagnóstico (definitivo)", "diagnostico"),
        (25, "2.2. Plan de Tratamiento inicial", "plan"),
        (26, "2.3. Pronóstico", "pronostico"),
        (27, "Cumple (SI/NO)", "cumple_summary")
    ]

    for r_num, title, *is_cat in rows_structure:
        ws.cell(row=r_num, column=1, value=title)
        ws.cell(row=r_num, column=1).alignment = left_aligned
        if is_cat and is_cat[0] is True:
            ws.cell(row=r_num, column=1).font = bold_font_navy
            ws.cell(row=r_num, column=1).fill = soft_teal_fill
        elif r_num == 27:
            ws.cell(row=r_num, column=1).font = bold_font_navy
            ws.cell(row=r_num, column=1).fill = soft_teal_fill

    ws.cell(row=28, column=1, value="Total Cumple").font = bold_font_navy
    ws.cell(row=28, column=1).fill = soft_teal_fill

    total_cumple = 0
    total_no_cumple = 0

    for idx in range(18):
        num_pauta = idx + 1
        col_start = 2 + (idx * 2)
        col_end = col_start + 1

        data = reg_dict.get(num_pauta) or {}

        ws.merge_cells(start_row=3, start_column=col_start, end_row=3, end_column=col_end)
        ws.cell(row=3, column=col_start, value=centro).alignment = center_aligned

        ws.merge_cells(start_row=4, start_column=col_start, end_row=4, end_column=col_end)
        ws.cell(row=4, column=col_start, value=ano).alignment = center_aligned

        ws.merge_cells(start_row=5, start_column=col_start, end_row=5, end_column=col_end)
        ws.cell(row=5, column=col_start, value=mes).alignment = center_aligned

        ws.merge_cells(start_row=6, start_column=col_start, end_row=6, end_column=col_end)
        ws.cell(row=6, column=col_start, value=grupo).alignment = center_aligned

        ws.merge_cells(start_row=7, start_column=col_start, end_row=7, end_column=col_end)
        ws.cell(row=7, column=col_start, value=especialidad).alignment = center_aligned

        keys_meta = ["rut_pac", "nom_pac", "fecha_diag", "nom_prof", "rut_prof"]
        for r_offset, k_m in enumerate(keys_meta, start=8):
            ws.merge_cells(start_row=r_offset, start_column=col_start, end_row=r_offset, end_column=col_end)
            ws.cell(row=r_offset, column=col_start, value=data.get(k_m, "")).alignment = center_aligned

        ws.merge_cells(start_row=13, start_column=col_start, end_row=13, end_column=col_end)
        ws.cell(row=13, column=col_start, value=num_pauta).alignment = center_aligned
        ws.cell(row=13, column=col_start).font = bold_font_navy
        ws.cell(row=13, column=col_start).fill = soft_teal_fill

        ws.cell(row=14, column=col_start, value="SI").alignment = center_aligned
        ws.cell(row=14, column=col_start).font = bold_font_navy
        ws.cell(row=14, column=col_end, value="NO").alignment = center_aligned
        ws.cell(row=14, column=col_end).font = bold_font_navy

        criteria_keys = [
            (17, "motivo"), (18, "patologias"), (19, "medicamentos"), (20, "alergias"),
            (21, "extraoral"), (22, "intraoral"), (24, "diagnostico"), (25, "plan"), (26, "pronostico")
        ]

        for r_i, k_i in criteria_keys:
            val_crit = data.get(k_i, "")
            if val_crit == "SI":
                ws.cell(row=r_i, column=col_start, value="√").alignment = center_aligned
            elif val_crit == "NO":
                ws.cell(row=r_i, column=col_end, value="X").alignment = center_aligned

        ws.merge_cells(start_row=15, start_column=col_start, end_row=15, end_column=col_end)
        ws.merge_cells(start_row=16, start_column=col_start, end_row=16, end_column=col_end)
        ws.merge_cells(start_row=23, start_column=col_start, end_row=23, end_column=col_end)

        c_gen = data.get("cumple", "")
        ws.merge_cells(start_row=27, start_column=col_start, end_row=27, end_column=col_end)
        if c_gen == "SI":
            ws.cell(row=27, column=col_start, value="SI").alignment = center_aligned
            total_cumple += 1
        elif c_gen == "NO":
            ws.cell(row=27, column=col_start, value="NO").alignment = center_aligned
            total_no_cumple += 1

    ws.merge_cells('B28:F28')
    ws['B28'] = total_cumple
    ws['B28'].alignment = center_aligned

    ws.merge_cells('G28:J28')
    ws['G28'] = "Total No Cumple"
    ws['G28'].font = bold_font_navy
    ws['G28'].alignment = center_aligned

    ws.merge_cells('K28:N28')
    ws['K28'] = total_no_cumple
    ws['K28'].alignment = center_aligned

    ws.merge_cells('O28:R28')
    ws['O28'] = "% Cumplimiento"
    ws['O28'].font = bold_font_navy
    ws['O28'].alignment = center_aligned

    completadas = sum(1 for v in reg_dict.values() if v is not None)
    porcentaje = f"{(total_cumple/18)*100:.1f}%" if completadas == 18 else f"{(total_cumple/completadas)*100:.1f}%" if completadas > 0 else "-"
    ws.merge_cells('S28:V28')
    ws['S28'] = porcentaje
    ws['S28'].alignment = center_aligned

    ws.merge_cells('A29:Q33')
    ws['A29'] = "Observaciones:"
    ws['A29'].font = bold_font_navy
    ws['A29'].alignment = Alignment(horizontal="left", vertical="top")

    ws.merge_cells('R29:U31')
    ws.merge_cells('R32:U33')
    ws['R32'] = "Nombre o Timbre\ndel responsable de\naplicar la pauta"
    ws['R32'].font = bold_font_navy
    ws['R32'].alignment = center_aligned

    ws.row_dimensions[32].height = 20
    ws.row_dimensions[33].height = 20

    for r in range(1, 29):
        for c in range(1, 38):
            ws.cell(row=r, column=c).border = thin_border

    for r in range(29, 34):
        for c in range(1, 22):
            ws.cell(row=r, column=c).border = thin_border

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


# ==============================================================================
# --- VISTA 1: PORTAL DE INICIO ---
# ==============================================================================
if st.session_state.pagina_activa == "inicio":
    st.title("RedSalud | Portal de Pautas de Supervisión")
    st.write("Bienvenido al sistema de consolidación de pautas clínicas. Seleccione la pauta a evaluar:")

    st.markdown("---")

    # Tarjeta 1: Pausa de Seguridad Dental
    st.markdown("""
    <div class="card-pauta">
        <h3 style="margin-top:0; color:#00205B;">🦷 Pausa de Seguridad Dental (GCL 2.1 AO)</h3>
        <p>Evaluación de Pausa de Seguridad Dental en Box Dental, Pabellón de Cirugía Menor e Imagenología.</p>
        <p><b>Formato:</b> Consolidado de 18 Pautas.</p>
    </div>
    """, unsafe_allow_html=True)
    
    if st.button("🚀 Ingresar a Pausa de Seguridad Dental", type="primary", use_container_width=True):
        st.session_state.pagina_activa = "pauta_dental"
        st.rerun()

    st.markdown("<br>", unsafe_allow_html=True)

    # Tarjeta 2: Higiene de Manos
    st.markdown("""
    <div class="card-pauta">
        <h3 style="margin-top:0; color:#00205B;">🧼 Higiene de Manos - Área Dental (GCL 1.2)</h3>
        <p>Supervisión del cumplimiento de técnica y momentos de Higiene de Manos en el equipo dental.</p>
        <p><b>Formato:</b> Consolidado de 12 Evaluaciones Mensuales.</p>
    </div>
    """, unsafe_allow_html=True)
    
    if st.button("🚀 Ingresar a Higiene de Manos Dental", type="primary", use_container_width=True):
        st.session_state.pagina_activa = "pauta_higiene"
        st.rerun()

    st.markdown("<br>", unsafe_allow_html=True)

    # Tarjeta 3: REG 1.2 AO
    st.markdown("""
    <div class="card-pauta">
        <h3 style="margin-top:0; color:#00205B;">📄 Registros Mínimos en Ficha Clínica Dental (REG 1.2 AO)</h3>
        <p>Supervisión de Anamnesis, Examen Físico, Diagnóstico, Plan de Tratamiento y Pronóstico en Ficha Clínica.</p>
        <p><b>Formato:</b> Consolidado de 18 Pautas con Carga Masiva desde Gema.</p>
    </div>
    """, unsafe_allow_html=True)
    
    if st.button("🚀 Ingresar a REG 1.2 AO (Registros Mínimos Ficha Clínica)", type="primary", use_container_width=True):
        st.session_state.pagina_activa = "pauta_reg12"
        st.rerun()


# ==============================================================================
# --- VISTA 2: PAUSA DE SEGURIDAD DENTAL (18 PAUTAS) ---
# ==============================================================================
elif st.session_state.pagina_activa == "pauta_dental":
    if st.button("⬅️ Volver al Portal de Inicio", use_container_width=True):
        st.session_state.pagina_activa = "inicio"
        st.rerun()

    st.markdown("---")
    st.title("RedSalud | Pausa de Seguridad Dental")

    with st.expander("📥 **Carga Masiva Mensual (Copiar y Pegar desde Excel / Archivo)**", expanded=False):
        st.write("Copia la tabla desde Excel y pégala abajo, o sube el archivo directamente:")
        fecha_sup_masiva = st.date_input("Fecha de Supervisión para el lote:", value=datetime.date.today())
        
        tab1, tab2 = st.tabs(["📋 Pegar Texto desde Excel", "📁 Subir Archivo Excel/CSV"])
        with tab1:
            texto_pegado = st.text_area("Pega la tabla copiada desde Excel aquí:", height=150)
            if st.button("⚡ Procesar Texto Pegado", type="primary"):
                if texto_pegado.strip():
                    try:
                        df_pasted = pd.read_csv(io.StringIO(texto_pegado), sep='\t')
                        cargadas = procesar_df_masivo(df_pasted, fecha_sup_masiva)
                        st.success(f"✅ ¡Se cargaron {cargadas} pautas automáticamente!")
                        st.rerun()
                    except Exception as e:
                        st.error("Ocurrió un error al leer el texto.")
                else:
                    st.warning("Pega texto antes de procesar.")
                    
        with tab2:
            archivo_subido = st.file_uploader("Selecciona archivo Excel (.xlsx) o CSV:", type=["xlsx", "csv"])
            if st.button("⚡ Procesar Archivo Subido", type="primary"):
                if archivo_subido is not None:
                    try:
                        df_file = pd.read_csv(archivo_subido) if archivo_subido.name.endswith('.csv') else pd.read_excel(archivo_subido)
                        cargadas = procesar_df_masivo(df_file, fecha_sup_masiva)
                        st.success(f"✅ ¡Se cargaron {cargadas} pautas automáticamente!")
                        st.rerun()
                    except Exception as e:
                        st.error(f"Error al leer archivo: {e}")

    st.markdown("---")
    completadas = sum(1 for v in st.session_state.pautas_data.values() if v is not None)
    st.progress(completadas / 18)
    st.caption(f"Progreso global: **{completadas} de 18 pautas guardadas**")

    pauta_seleccionada = st.selectbox(
        "Selecciona la pauta a ingresar o revisar:",
        options=list(range(1, 19)),
        index=st.session_state.pauta_actual - 1,
        format_func=lambda num: f"Pauta N° {num} ({'✅ Guardada' if st.session_state.pautas_data[num] is not None else '⏳ Pendiente'})"
    )

    st.session_state.pauta_actual = pauta_seleccionada
    p_num = st.session_state.pauta_actual
    datos_existentes = st.session_state.pautas_data[p_num] or {}

    servicios = ["Sala de Procedimiento Dental (BD)", "Pabellón de Cirugía menor Dental (PD)", "Imagenología Dental (RX)"]
    idx_serv = servicios.index(datos_existentes.get('servicio')) if datos_existentes.get('servicio') in servicios else 0
    idx_exo = 0 if datos_existentes.get('exodoncia') != "NO" else 1
    idx_cumple = 0 if datos_existentes.get('cumple') != "NO" else 1

    st.markdown('<div class="card-pauta">', unsafe_allow_html=True)
    st.subheader(f"Formulario Pauta N° {p_num}")

    centro = st.text_input("Centro", value=datos_existentes.get('centro', st.session_state.ultimo_centro), key=f"c_{p_num}")
    fecha_sup = st.date_input("Fecha de Supervisión", value=parse_fecha(datos_existentes.get('fecha_sup')), key=f"fs_{p_num}")

    nombre = st.text_input("Nombre de la persona supervisada", value=datos_existentes.get('nombre', ''), key=f"n_{p_num}")
    apellido = st.text_input("Apellido(s) de la persona supervisada", value=datos_existentes.get('apellido', ''), key=f"a_{p_num}")

    rut = st.text_input("RUT del paciente", value=datos_existentes.get('rut', ''), key=f"r_{p_num}")
    fecha_atencion = st.date_input("Fecha de Atención supervisada", value=parse_fecha(datos_existentes.get('fecha_atencion')), key=f"fa_{p_num}")

    servicio = st.selectbox("Servicio Clínico", servicios, index=idx_serv, key=f"s_{p_num}")
    exodoncia = st.radio("¿Procedimiento corresponde a Exodoncia?", ["SI", "NO"], index=idx_exo, key=f"e_{p_num}")

    st.markdown("---")
    st.markdown("**CRITERIO A EVALUAR**")
    cumple = st.radio("¿Se constata Pausa de Seguridad Dental realizada y registrada en Ficha Clínica?", ["SI", "NO"], index=idx_cumple, key=f"cu_{p_num}")

    if st.button(f"💾 Guardar Pauta N° {p_num}", type="primary", use_container_width=True, key=f"btn_{p_num}"):
        st.session_state.pautas_data[p_num] = {
            "centro": clean_text_spaces(centro),
            "fecha_sup": fecha_sup.strftime("%d-%m-%Y"),
            "nombre": clean_text_spaces(nombre),
            "apellido": clean_text_spaces(apellido),
            "rut": re.sub(r'[^0-9kK\-]', '', rut).strip().upper(),
            "fecha_atencion": fecha_atencion.strftime("%d-%m-%Y"),
            "servicio": servicio,
            "exodoncia": exodoncia,
            "cumple": cumple
        }
        st.session_state.ultimo_centro = clean_text_spaces(centro)
        if p_num < 18:
            st.session_state.pauta_actual = p_num + 1
        st.rerun()

    st.markdown('</div>', unsafe_allow_html=True)

    excel_file = generar_excel_pausa_dental(st.session_state.pautas_data)
    st.download_button(label="📥 Descargar Excel Consolidado Pausa Dental", data=excel_file, file_name="Consolidado_Pausa_Seguridad_Dental.xlsx", mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", use_container_width=True)

    if st.button("🔄 Reiniciar todo y borrar pautas", use_container_width=True):
        st.session_state.pautas_data = {i: None for i in range(1, 19)}
        st.session_state.ultimo_centro = ""
        st.session_state.pauta_actual = 1
        st.rerun()

    pautas_list = [v for v in st.session_state.pautas_data.values() if v is not None]
    if len(pautas_list) > 0:
        with st.expander(f"📋 Ver resumen de pautas guardadas ({len(pautas_list)}/18)", expanded=True):
            st.dataframe(pd.DataFrame(pautas_list), use_container_width=True)


# ==============================================================================
# --- VISTA 3: HIGIENE DE MANOS - ÁREA DENTAL (12 EVALUACIONES) ---
# ==============================================================================
elif st.session_state.pagina_activa == "pauta_higiene":
    if st.button("⬅️ Volver al Portal de Inicio", use_container_width=True):
        st.session_state.pagina_activa = "inicio"
        st.rerun()

    st.markdown("---")
    st.title("RedSalud | Higiene de Manos - Área Dental (GCL 1.2)")

    col1, col2 = st.columns(2)
    with col1:
        st.session_state.centro_higiene = clean_text_spaces(st.text_input("Centro Dental", value=st.session_state.centro_higiene))
    with col2:
        st.session_state.mes_higiene = clean_text_spaces(st.text_input("Mes de Evaluación", value=st.session_state.mes_higiene))

    st.session_state.responsable_higiene = clean_text_spaces(st.text_input("Nombre de responsable del indicador", value=st.session_state.responsable_higiene))

    st.markdown("---")
    completadas_h = sum(1 for v in st.session_state.higiene_data.values() if v is not None)
    st.progress(completadas_h / 12)
    st.caption(f"Progreso global: **{completadas_h} de 12 evaluaciones guardadas**")

    h_num = st.selectbox(
        "Selecciona la evaluación a ingresar o revisar:",
        options=list(range(1, 13)),
        index=st.session_state.higiene_actual - 1,
        format_func=lambda num: f"Evaluación N° {num} ({'✅ Guardada' if st.session_state.higiene_data[num] is not None else '⏳ Pendiente'})"
    )

    st.session_state.higiene_actual = h_num
    datos_h = st.session_state.higiene_data[h_num] or {}

    st.markdown('<div class="card-pauta">', unsafe_allow_html=True)
    st.subheader(f"Evaluación N° {h_num}")

    fecha_eval = st.date_input("Fecha de la evaluación", value=parse_fecha(datos_h.get('fecha_eval')), key=f"feh_{h_num}")
    evaluador = st.text_input("Nombre del evaluador", value=datos_h.get('evaluador', st.session_state.evaluador_higiene), key=f"evr_{h_num}")
    evaluado = st.text_input("Nombre del evaluado", value=datos_h.get('evaluado', ''), key=f"evd_{h_num}")

    st.markdown("---")
    st.markdown("**CRITERIOS A EVALUAR**")

    oportunidades_map = {
        "1": "1. Antes del contacto con el paciente",
        "2": "2. Antes de una técnica aséptica",
        "3": "3. Después de la exposición a fluidos corporales o manejo de fluidos contaminados",
        "4": "4. Después del contacto con el paciente",
        "5": "5. Después de tener contacto con la zona alrededor del paciente"
    }

    op_val = str(datos_h.get('oportunidad', '1'))
    op_sel = st.selectbox("N° Oportunidad evaluada", options=["1", "2", "3", "4", "5"], index=int(op_val)-1 if op_val in ["1","2","3","4","5"] else 0, format_func=lambda x: oportunidades_map[x], key=f"op_{h_num}")

    idx_cumple_h = 0 if datos_h.get('cumple') != "NO" else 1
    cumple_h = st.radio("Se realiza higiene de manos según oportunidad", ["SI", "NO"], index=idx_cumple_h, key=f"ch_{h_num}")

    if st.button(f"💾 Guardar Evaluación N° {h_num}", type="primary", use_container_width=True, key=f"btn_h_{h_num}"):
        st.session_state.higiene_data[h_num] = {
            "fecha_eval": fecha_eval.strftime("%d/%m"),
            "evaluador": clean_text_spaces(evaluador),
            "evaluado": clean_text_spaces(evaluado),
            "oportunidad": op_sel,
            "cumple": cumple_h
        }
        st.session_state.evaluador_higiene = clean_text_spaces(evaluador)
        if h_num < 12:
            st.session_state.higiene_actual = h_num + 1
        st.rerun()

    st.markdown('</div>', unsafe_allow_html=True)

    excel_higiene = generar_excel_higiene_manos(
        st.session_state.higiene_data,
        st.session_state.centro_higiene,
        st.session_state.mes_higiene,
        st.session_state.responsable_higiene
    )

    st.download_button(
        label="📥 Descargar Excel Consolidado Higiene de Manos (GCL 1.2)",
        data=excel_higiene,
        file_name="Consolidado_Higiene_de_Manos_Dental.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        use_container_width=True
    )

    if st.button("🔄 Reiniciar todo y borrar evaluaciones", use_container_width=True):
        st.session_state.higiene_data = {i: None for i in range(1, 13)}
        st.session_state.higiene_actual = 1
        st.rerun()

    higiene_list = [v for v in st.session_state.higiene_data.values() if v is not None]
    if len(higiene_list) > 0:
        with st.expander(f"📋 Ver resumen de evaluaciones guardadas ({len(higiene_list)}/12)", expanded=True):
            st.dataframe(pd.DataFrame(higiene_list), use_container_width=True)


# ==============================================================================
# --- VISTA 4: REGISTROS MÍNIMOS EN FICHA CLÍNICA DENTAL (REG 1.2 AO) ---
# ==============================================================================
elif st.session_state.pagina_activa == "pauta_reg12":
    if st.button("⬅️ Volver al Portal de Inicio", use_container_width=True):
        st.session_state.pagina_activa = "inicio"
        st.rerun()

    st.markdown("---")
    st.title("RedSalud | Registros Mínimos en Ficha Clínica Dental (REG 1.2 AO)")

    # Datos Generales
    c1, c2, c3 = st.columns(3)
    with c1:
        st.session_state.centro_reg12 = clean_text_spaces(st.text_input("Centro", value=st.session_state.centro_reg12))
        st.session_state.grupo_reg12 = clean_text_spaces(st.text_input("Grupo", value=st.session_state.grupo_reg12))
    with c2:
        st.session_state.ano_reg12 = clean_text_spaces(st.text_input("Año", value=st.session_state.ano_reg12))
        st.session_state.especialidad_reg12 = clean_text_spaces(st.text_input("Especialidad", value=st.session_state.especialidad_reg12))
    with c3:
        st.session_state.mes_reg12 = clean_text_spaces(st.text_input("Mes", value=st.session_state.mes_reg12))

    # Carga Masiva desde Gema
    with st.expander("📥 **Carga Masiva desde Gema / Excel (Copiar y Pegar)**", expanded=False):
        st.write("Pega el cuadro completo generado por la gema aquí abajo:")
        texto_gema = st.text_area("Pega la tabla de la gema aquí:", height=180, placeholder="RUT PACIENTENOMBRE PACIENTEFECHA DIAGNÓSTICO...")
        if st.button("⚡ Procesar Cuadro de la Gema", type="primary"):
            if texto_gema.strip():
                cargadas_r = procesar_texto_reg12(texto_gema)
                st.success(f"✅ ¡Se cargaron {cargadas_r} pautas de REG 1.2 AO automáticamente!")
                st.rerun()
            else:
                st.warning("Pega el texto antes de procesar.")

    st.markdown("---")
    completadas_r = sum(1 for v in st.session_state.reg12_data.values() if v is not None)
    st.progress(completadas_r / 18)
    st.caption(f"Progreso global: **{completadas_r} de 18 pautas guardadas**")

    r_num = st.selectbox(
        "Selecciona la pauta a ingresar o revisar:",
        options=list(range(1, 19)),
        index=st.session_state.reg12_actual - 1,
        format_func=lambda num: f"Pauta N° {num} ({'✅ Guardada' if st.session_state.reg12_data[num] is not None else '⏳ Pendiente'})"
    )

    st.session_state.reg12_actual = r_num
    datos_r = st.session_state.reg12_data[r_num] or {}

    st.markdown('<div class="card-pauta">', unsafe_allow_html=True)
    st.subheader(f"Formulario Pauta N° {r_num}")

    col_a, col_b = st.columns(2)
    with col_a:
        rut_pac = st.text_input("RUT Paciente", value=datos_r.get('rut_pac', ''), key=f"rp_{r_num}")
        nom_pac = st.text_input("Nombre Paciente", value=datos_r.get('nom_pac', ''), key=f"np_{r_num}")
        fecha_diag = st.date_input("Fecha Diagnóstico", value=parse_fecha(datos_r.get('fecha_diag')), key=f"fd_{r_num}")
    with col_b:
        nom_prof = st.text_input("Nombre Profesional", value=datos_r.get('nom_prof', ''), key=f"npr_{r_num}")
        rut_prof = st.text_input("RUT Profesional", value=datos_r.get('rut_prof', ''), key=f"rpr_{r_num}")

    st.markdown("---")
    st.markdown("**1. HISTORIA CLÍNICA**")
    st.markdown("*1.1 Anamnesis*")
    
    m_motivo = st.radio("Motivo de Consulta", ["SI", "NO"], index=0 if datos_r.get('motivo') != "NO" else 1, key=f"r_mot_{r_num}")
    m_pat = st.radio("Patologías", ["SI", "NO"], index=0 if datos_r.get('patologias') != "NO" else 1, key=f"r_pat_{r_num}")
    m_med = st.radio("Medicamentos", ["SI", "NO"], index=0 if datos_r.get('medicamentos') != "NO" else 1, key=f"r_med_{r_num}")
    m_ale = st.radio("Alergias", ["SI", "NO"], index=0 if datos_r.get('alergias') != "NO" else 1, key=f"r_ale_{r_num}")

    st.markdown("*Exámenes Fisicos*")
    m_ext = st.radio("1.2. Examen Extraoral", ["SI", "NO"], index=0 if datos_r.get('extraoral') != "NO" else 1, key=f"r_ext_{r_num}")
    m_int = st.radio("1.3. Examen Intraoral", ["SI", "NO"], index=0 if datos_r.get('intraoral') != "NO" else 1, key=f"r_int_{r_num}")

    st.markdown("---")
    st.markdown("**2. DIAGNÓSTICO Y PLAN DE TRATAMIENTO**")
    m_dia = st.radio("2.1. Diagnóstico (definitivo)", ["SI", "NO"], index=0 if datos_r.get('diagnostico') != "NO" else 1, key=f"r_dia_{r_num}")
    m_pla = st.radio("2.2. Plan de Tratamiento inicial", ["SI", "NO"], index=0 if datos_r.get('plan') != "NO" else 1, key=f"r_pla_{r_num}")
    m_pro = st.radio("2.3. Pronóstico", ["SI", "NO"], index=0 if datos_r.get('pronostico') != "NO" else 1, key=f"r_pro_{r_num}")

    st.markdown("---")
    m_cum = st.radio("Cumple (SI/NO)", ["SI", "NO"], index=0 if datos_r.get('cumple') != "NO" else 1, key=f"r_cum_{r_num}")

    if st.button(f"💾 Guardar Pauta REG 1.2 N° {r_num}", type="primary", use_container_width=True, key=f"btn_reg_{r_num}"):
        st.session_state.reg12_data[r_num] = {
            "rut_pac": re.sub(r'[^0-9kK\-]', '', rut_pac).strip().upper(),
            "nom_pac": clean_text_spaces(nom_pac),
            "fecha_diag": fecha_diag.strftime("%d-%m-%Y"),
            "nom_prof": clean_text_spaces(nom_prof),
            "rut_prof": re.sub(r'[^0-9kK\-]', '', rut_prof).strip().upper(),
            "motivo": m_motivo,
            "patologias": m_pat,
            "medicamentos": m_med,
            "alergias": m_ale,
            "extraoral": m_ext,
            "intraoral": m_int,
            "diagnostico": m_dia,
            "plan": m_pla,
            "pronostico": m_pro,
            "cumple": m_cum
        }
        if r_num < 18:
            st.session_state.reg12_actual = r_num + 1
        st.rerun()

    st.markdown('</div>', unsafe_allow_html=True)

    excel_reg12 = generar_excel_reg12(
        st.session_state.reg12_data,
        st.session_state.centro_reg12,
        st.session_state.ano_reg12,
        st.session_state.mes_reg12,
        st.session_state.grupo_reg12,
        st.session_state.especialidad_reg12
    )

    st.download_button(
        label="📥 Descargar Excel Consolidado REG 1.2 AO",
        data=excel_reg12,
        file_name="Consolidado_Registros_Minimos_Ficha_Clinica_REG12.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        use_container_width=True
    )

    if st.button("🔄 Reiniciar todo y borrar pautas REG 1.2", use_container_width=True):
        st.session_state.reg12_data = {i: None for i in range(1, 19)}
        st.session_state.reg12_actual = 1
        st.rerun()

    reg12_list = [v for v in st.session_state.reg12_data.values() if v is not None]
    if len(reg12_list) > 0:
        with st.expander(f"📋 Ver resumen de pautas guardadas ({len(reg12_list)}/18)", expanded=True):
            st.dataframe(pd.DataFrame(reg12_list), use_container_width=True)
